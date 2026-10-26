import shutil
import subprocess
from pathlib import Path

from dotenv import load_dotenv
from langchain_community.document_loaders import PyMuPDFLoader
from langchain_community.vectorstores import FAISS
from langchain_groq import ChatGroq
from langchain_community.embeddings import FastEmbedEmbeddings
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from ocr import process_pdfs, load_cache

load_dotenv()

DOCS_DIR = Path("docs")
INDEX_PATH = Path("faiss_index")
EMBED_MODEL = "BAAI/bge-small-en-v1.5"

PROMPT = """\
Use the context below to answer the question. Be concise and direct.
If the answer is not present in the context, say "I don't have enough information to answer that."

Context:
{context}

Question: {question}
Answer:"""


def load_pdfs():
    pdf_files = sorted(DOCS_DIR.glob("*.pdf"))
    print(f"Found {len(pdf_files)} PDF(s).")

    print("Checking PDFs for OCR requirements...")
    process_pdfs(pdf_files)
    cache = load_cache()

    docs = []
    for path in pdf_files:
        name = path.name
        print(f"  Reading {name}...")
        entry = cache.get(name, {})

        if entry.get("type") == "ocr":
            pages = entry.get("pages", [])
            for i, page_text in enumerate(pages):
                if page_text.strip():
                    docs.append(Document(
                        page_content=page_text.strip(),
                        metadata={"source": str(path), "page": i}
                    ))
        else:
            loader = PyMuPDFLoader(str(path))
            docs.extend(loader.load())

    return docs


def _convert_ppt_to_pptx(ppt_path):
    libreoffice = shutil.which("libreoffice") or shutil.which("soffice")
    if not libreoffice:
        print(f"  {ppt_path.name}: .ppt format requires LibreOffice. Install it or convert to .pptx manually.")
        return None

    print(f"  {ppt_path.name}: converting via LibreOffice...")
    subprocess.run(
        [libreoffice, "--headless", "--convert-to", "pptx", "--outdir", str(ppt_path.parent), str(ppt_path)],
        capture_output=True
    )

    converted = ppt_path.with_suffix(".pptx")
    if converted.exists():
        return converted

    print(f"  {ppt_path.name}: LibreOffice conversion failed.")
    return None


def _open_presentation(path):
    from pptx import Presentation

    try:
        return Presentation(str(path))
    except Exception:
        if path.suffix.lower() == ".ppt":
            converted = _convert_ppt_to_pptx(path)
            if converted:
                try:
                    return Presentation(str(converted))
                except Exception as e:
                    print(f"  {path.name}: could not open after conversion: {e}")
        else:
            print(f"  {path.name}: could not open.")
        return None


def load_pptx_docs():
    try:
        from pptx import Presentation
    except ImportError:
        raise ImportError("python-pptx is required. Run: pip install python-pptx")

    docs = []
    ppt_files = sorted(list(DOCS_DIR.glob("*.pptx")) + list(DOCS_DIR.glob("*.ppt")))
    print(f"Found {len(ppt_files)} PPT/PPTX file(s).")

    for path in ppt_files:
        print(f"  Reading {path.name}...")
        prs = _open_presentation(path)
        if prs is None:
            continue

        skipped = 0
        for i, slide in enumerate(prs.slides):
            text_parts = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = " ".join(run.text for run in para.runs).strip()
                        if line:
                            text_parts.append(line)

                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(
                            cell.text.strip()
                            for cell in row.cells
                            if cell.text.strip()
                        )
                        if row_text:
                            text_parts.append(row_text)

            slide_text = "\n".join(text_parts).strip()
            if not slide_text:
                skipped += 1
                continue

            docs.append(Document(
                page_content=slide_text,
                metadata={"source": str(path), "page": i}
            ))

        if skipped:
            print(f"  {path.name}: {skipped} image-only slide(s) skipped.")

    return docs


def load_docs():
    docs = []
    docs.extend(load_pdfs())
    docs.extend(load_pptx_docs())
    return docs


def build_index():
    print("Loading documents...")
    docs = load_docs()
    if not docs:
        raise ValueError("No documents found in the docs/ folder. Add .pdf, .pptx, or .ppt files.")

    print(f"Loaded {len(docs)} page(s) across all files. Chunking...")
    splitter = RecursiveCharacterTextSplitter(chunk_size=500, chunk_overlap=50)
    chunks = splitter.split_documents(docs)
    print(f"Created {len(chunks)} chunk(s). Generating embeddings...")

    embeddings = FastEmbedEmbeddings(model_name=EMBED_MODEL)
    index = FAISS.from_documents(chunks, embeddings)
    index.save_local(str(INDEX_PATH))
    print("Index saved.")
    return index


def load_index():
    embeddings = FastEmbedEmbeddings(model_name=EMBED_MODEL)
    return FAISS.load_local(
        str(INDEX_PATH), embeddings, allow_dangerous_deserialization=True
    )


def get_index():
    if INDEX_PATH.exists():
        print("Loading existing index...")
        return load_index()
    print("No index found. Building from scratch...")
    return build_index()


def retrieve(query, index, k=3):
    return index.similarity_search(query, k=k)


def answer_query(query, chunks):
    context = "\n\n---\n\n".join(chunk.page_content for chunk in chunks)
    prompt = PROMPT.format(context=context, question=query)
    llm = ChatGroq(model="openai/gpt-oss-120b", temperature=0.1)
    return llm.invoke(prompt).content